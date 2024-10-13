# -*- coding: utf-8 -*-
"""Módulo para processamento de arquivos com integração ao GPT-4 e banco de dados SQLite."""

import pandas as pd
import PyPDF2
import io
from openpyxl import load_workbook
from docx import Document
from datetime import datetime
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app_balance.services.openai_service import analyze_data  # Integração com GPT-4
from app_balance.models import Prompt, Recebimento  # Modelos para salvar prompts e recebimentos

# Configuração do SQLite
DATABASE_URL = 'sqlite:///recebimentos.db'

# Configurar a engine e o session para o SQLite
engine = create_engine(DATABASE_URL)
Session = sessionmaker(bind=engine)
session = Session()

def processar_arquivo_excel(arquivo: bytes) -> str:
    """Processa um arquivo Excel (.xlsx), extrai recebimentos, salva no DB e envia análise ao GPT-4."""
    try:
        with io.BytesIO(arquivo) as excel_file:
            wb = load_workbook(excel_file)
            sheet = wb.active
            df = pd.DataFrame(sheet.values)
        
        recebimentos = []
        for _, row in df.iterrows():
            try:
                recebimento = Recebimento(
                    data=datetime.strptime(row[0], '%Y-%m-%d').date(),
                    valor=float(row[1])
                )
                recebimentos.append(recebimento)
            except (ValueError, IndexError):
                continue  # Ignora linhas com erros de conversão
        
        salvar_recebimentos_no_db(recebimentos)
        
        # Gerar prompt para GPT-4
        prompt = gerar_prompt_gpt(recebimentos)
        salvar_prompt_no_db(prompt)
        
        # Enviar ao GPT-4
        resposta_gpt = analyze_data(prompt)
        return resposta_gpt
    
    except Exception as e:
        raise RuntimeError(f"Erro ao processar arquivo Excel: {str(e)}")


def processar_arquivo_pdf(arquivo: bytes) -> str:
    """Processa um arquivo PDF, extrai recebimentos, salva no DB e envia análise ao GPT-4."""
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(arquivo))
        recebimentos = []
        for page in reader.pages:
            texto = page.extract_text()
            linhas = texto.split('\n')
            for linha in linhas:
                try:
                    data_str, valor_str = linha.split(' ')
                    recebimento = Recebimento(
                        data=datetime.strptime(data_str, '%Y-%m-%d').date(),
                        valor=float(valor_str)
                    )
                    recebimentos.append(recebimento)
                except (ValueError, IndexError):
                    continue  # Ignora linhas com erros de conversão
        
        salvar_recebimentos_no_db(recebimentos)
        
        # Gerar prompt para GPT-4
        prompt = gerar_prompt_gpt(recebimentos)
        salvar_prompt_no_db(prompt)
        
        # Enviar ao GPT-4
        resposta_gpt = analyze_data(prompt)
        return resposta_gpt
    
    except Exception as e:
        raise RuntimeError(f"Erro ao processar arquivo PDF: {str(e)}")


def processar_arquivo_docx(arquivo: bytes) -> str:
    """Processa um arquivo DOCX, extrai recebimentos, salva no DB e envia análise ao GPT-4."""
    try:
        documento = Document(io.BytesIO(arquivo))
        recebimentos = []
        for para in documento.paragraphs:
            try:
                data_str, valor_str = para.text.split(' ')
                recebimento = Recebimento(
                    data=datetime.strptime(data_str, '%Y-%m-%d').date(),
                    valor=float(valor_str)
                )
                recebimentos.append(recebimento)
            except (ValueError, IndexError):
                continue  # Ignora parágrafos com erros de conversão
        
        salvar_recebimentos_no_db(recebimentos)
        
        # Gerar prompt para GPT-4
        prompt = gerar_prompt_gpt(recebimentos)
        salvar_prompt_no_db(prompt)
        
        # Enviar ao GPT-4
        resposta_gpt = analyze_data(prompt)
        return resposta_gpt
    
    except Exception as e:
        raise RuntimeError(f"Erro ao processar arquivo DOCX: {str(e)}")


def salvar_recebimentos_no_db(recebimentos: list):
    """Salva uma lista de recebimentos no banco de dados SQLite."""
    try:
        session.add_all(recebimentos)
        session.commit()
    except Exception as e:
        session.rollback()
        raise RuntimeError(f"Erro ao salvar os recebimentos no banco de dados: {str(e)}")


def gerar_prompt_gpt(recebimentos: list) -> str:
    """
    Gera um prompt detalhado com base nos recebimentos extraídos para análise via GPT-4.
    
    Args:
        recebimentos (list): Lista de objetos Recebimento contendo dados de recebimentos.
    
    Returns:
        str: Prompt formatado para enviar ao GPT-4.
    """
    total_recebimentos = sum([r.valor for r in recebimentos])
    detalhes_recebimentos = "\n".join([f"Recebimento em {r.data}: R$ {r.valor:.2f}" for r in recebimentos])
    
    prompt = (
        f"Você recebeu os seguintes recebimentos ao longo do mês:\n{detalhes_recebimentos}\n\n"
        f"O total de recebimentos foi de R$ {total_recebimentos:.2f}.\n"
        "Por favor, forneça uma análise detalhada desses recebimentos, incluindo sugestões sobre como melhor gerenciá-los."
    )
    
    return prompt


def salvar_prompt_no_db(prompt: str):
    """Salva o prompt enviado ao GPT-4 no banco de dados."""
    try:
        novo_prompt = Prompt(conteudo=prompt, data=datetime.now())
        session.add(novo_prompt)
        session.commit()
    except Exception as e:
        session.rollback()
        raise RuntimeError(f"Erro ao salvar o prompt no banco de dados: {str(e)}")


def processar_arquivo(arquivo: bytes, tipo_arquivo: str) -> str:
    """Processa um arquivo com base no tipo fornecido (Excel, PDF ou DOCX) e realiza a análise prioritária via GPT-4."""
    try:
        if tipo_arquivo == 'excel':
            resposta_gpt = processar_arquivo_excel(arquivo)
        elif tipo_arquivo == 'pdf':
            resposta_gpt = processar_arquivo_pdf(arquivo)
        elif tipo_arquivo == 'docx':
            resposta_gpt = processar_arquivo_docx(arquivo)
        else:
            raise ValueError(f"Tipo de arquivo não suportado: {tipo_arquivo}")

        return f"Análise do GPT-4 concluída com sucesso:\n{resposta_gpt}"
    
    except Exception as e:
        raise RuntimeError(f"Erro ao processar arquivo: {str(e)}")
